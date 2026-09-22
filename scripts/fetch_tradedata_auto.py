# -*- coding: utf-8 -*-
"""
관세청 수출입무역통계(tradedata.go.kr) 10일 잠정치 자동 수집기
- 웹 브라우저 방문 없이 tradedata.go.kr 서버에서 직접 최신 잠정치 데이터 조회
- 엑셀 파일(.xlsx)로 자동 생성하여 'data/trade' 폴더에 저장
- 대시보드(trade_10days.html) 자동 갱신 트리거
"""

import os
import sys
import datetime
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DATA_DIR = os.path.join(BASE_DIR, "data", "trade")

sys.path.insert(0, os.path.join(BASE_DIR, "scripts"))
import update_trade_dashboard

HEADERS_MAP = [
    ("월별", "priodMon"),
    ("기간", "priodDt"),
    ("전체", "itemUsdAmt00"),
    ("반도체", "itemUsdAmt01"),
    ("철강제품", "itemUsdAmt02"),
    ("승용차", "itemUsdAmt03"),
    ("석유제품", "itemUsdAmt04"),
    ("무선통신기기", "itemUsdAmt05"),
    ("선박", "itemUsdAmt06"),
    ("자동차부품", "itemUsdAmt07"),
    ("컴퓨터주변기기", "itemUsdAmt08"),
    ("정밀기기", "itemUsdAmt09"),
    ("가전제품", "itemUsdAmt10")
]

def fetch_latest_tradedata(start_ym=None, end_ym=None):
    print(">> [1/4] 관세청 무역통계(tradedata.go.kr) 서버에 접속 중...", flush=True)
    session = requests.Session()
    session.headers.update({
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Accept': 'application/json, text/javascript, */*; q=0.01',
        'X-Requested-With': 'XMLHttpRequest',
        'Referer': 'https://tradedata.go.kr/cts/hmpg/openETS0100173Q.do?menuId=ETS_MNK_10500000'
    })

    # Initialize session
    try:
        session.get('https://tradedata.go.kr/cts/index.do', verify=False, timeout=10)
        session.get('https://tradedata.go.kr/cts/hmpg/openETS0100173Q.do?menuId=ETS_MNK_10500000', verify=False, timeout=10)
    except Exception as e:
        print(f"[경고] 세션 초기화 중 예외 (계속 진행): {e}", flush=True)

    # 1. Check maxYear
    max_year = "202609"
    try:
        r_box = session.post('https://tradedata.go.kr/cts/hmpg/retrieveSetSelectBoxTentative.do', data={'statsKind': 'P'}, verify=False, timeout=10)
        j_box = r_box.json()
        if 'item' in j_box and j_box['item'].get('maxYear'):
            max_year = j_box['item']['maxYear']
    except Exception as e:
        print(f"[알림] maxYear 조회 기본값 사용 ({max_year}): {e}", flush=True)

    if not end_ym:
        end_ym = max_year
    if not start_ym:
        # 최근 3개월간 조회 (예: 202609 기준 202607 ~ 202609, 총 3개월)
        y = int(end_ym[:4])
        m = int(end_ym[4:6])
        m -= 2
        if m <= 0:
            m += 12
            y -= 1
        start_ym = f"{y}{m:02d}"

    print(f">> [2/4] 최신 잠정치 데이터 3개월간 조회 (기간: {start_ym} ~ {end_ym})...", flush=True)

    params = {
        'statsKind': 'ETS_MNK_1050000A', # 품목별
        'imexTpcd': 'E',                 # 수출
        'priodKind': 'MON',
        'priodFr': start_ym,
        'priodTo': end_ym,
        'priodDate': '',                 # 전체 구간
        'selectPaging': '1',
        'showPagingLine': '100',
        'sortColumn': '',
        'sortOrder': ''
    }

    try:
        r_data = session.post('https://tradedata.go.kr/cts/hmpg/retrieveTentativeValues.do', data=params, verify=False, timeout=15)
        j_data = r_data.json()
    except Exception as e:
        print(f"[오류] 데이터 요청 실패: {e}", flush=True)
        return False

    if not j_data or 'items' not in j_data or len(j_data['items']) == 0:
        print("[오류] 데이터를 가져오지 못했습니다:", j_data, flush=True)
        return False

    raw_items = j_data['items']
    print(f">> [성공] 관세청 서버로부터 총 {len(raw_items)}개 잠정치 행을 수신했습니다.", flush=True)

    # 2. Build Excel File
    print(">> [3/4] 엑셀(.xlsx) 파일 생성 및 데이터 포맷팅 중...", flush=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "잠정치 통계(품목별)"

    # Meta headers
    ws.cell(1, 1, "잠정치 통계(품목별)")
    ws.cell(3, 1, f"검색조건  통계항목 : 품목별, 수출입구분 : 수출, 조회기간 : {start_ym} ~ {end_ym}, 전체")
    ws.cell(4, len(HEADERS_MAP), "단위:천 달러")

    # Column header
    for col_idx, (col_name, _) in enumerate(HEADERS_MAP, start=1):
        c = ws.cell(5, col_idx, col_name)
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")

    # Rows
    for row_idx, item in enumerate(raw_items, start=6):
        for col_idx, (_, key) in enumerate(HEADERS_MAP, start=1):
            val = item.get(key, "")
            if col_idx >= 3:
                # Number formatting
                try:
                    num_val = float(str(val).replace(",", "").strip())
                    ws.cell(row_idx, col_idx, num_val)
                    ws.cell(row_idx, col_idx).number_format = "#,##0"
                except ValueError:
                    ws.cell(row_idx, col_idx, val)
            else:
                ws.cell(row_idx, col_idx, val)

    today_str = datetime.datetime.now().strftime("%Y%m%d")
    out_excel_filename = f"잠정치 통계(품목별)_{today_str}.xlsx"
    out_excel_path = os.path.join(DATA_DIR, out_excel_filename)

    os.makedirs(DATA_DIR, exist_ok=True)
    wb.save(out_excel_path)
    print(f">> [저장 완료] 엑셀 파일이 저장되었습니다: {out_excel_path}", flush=True)

    # 3. Trigger Dashboard Regeneration
    print(">> [4/4] 대시보드 HTML 자동 갱신 실행 중...", flush=True)
    update_trade_dashboard.main()
    print(">> [완전 완료] 관세청 최신 잠정치 데이터 다운로드 및 대시보드 갱신이 성공적으로 끝났습니다!", flush=True)
    return True

if __name__ == "__main__":
    fetch_latest_tradedata()
